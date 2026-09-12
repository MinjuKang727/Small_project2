from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import sys

# ==========================================
# 설정 영역
# ==========================================
BASE_DIR = Path(__file__).parent
TXT_FILENAME = "Day10.txt"               # 읽어올 단어장 파일 이름
OUTPUT_PPTX_NAME = f"{Path(TXT_FILENAME).stem}_steps.pptx"

TXT_PATH = BASE_DIR / "txt" / TXT_FILENAME
OUTPUT_PPTX_PATH = BASE_DIR / "save" / OUTPUT_PPTX_NAME


def load_vocab_file(file_path):
    if not file_path.exists():
        print(f"\n❌ [오류] 단어장 파일을 찾을 수 없습니다: {file_path.resolve()}")
        sys.exit(1)
        
    vocab_list = []
    lines = file_path.read_text(encoding="utf-8").splitlines()
    for line in lines:
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        vocab_list.append(parse_line(line))
        
    if not vocab_list:
        print(f"\n❌ [오류] 파일({file_path.name}) 안에 읽어올 단어 데이터가 없습니다.")
        sys.exit(1)
        
    return vocab_list

def parse_line(raw_line):
    if "#" in raw_line:
        main_part, grammar = raw_line.split("#", 1)
        grammar_note = grammar.strip()
    else:
        main_part = raw_line
        grammar_note = ""

    parts = [p.strip() for p in main_part.split('/')]
    
    word = parts[0] if len(parts) > 0 else ""
    raw_meanings_block = parts[1] if len(parts) > 1 else ""
    
    meaning_segments = [seg.strip() for seg in raw_meanings_block.split("|") if seg.strip()]
    
    meaning_pairs = []  # (뜻, 동의어) 세트를 순서대로 저장
    extracted_similar = []
    extracted_antonyms = []
    extracted_derivatives = []

    for seg in meaning_segments:
        if "=" in seg:
            sub_p = seg.split("=", 1)
            m_text = sub_p[0].strip()
            s_text = sub_p[1].strip()
            meaning_pairs.append({"meaning": m_text, "synonym": s_text})
        else:
            meaning_pairs.append({"meaning": seg, "synonym": ""})

    # 슬래시 뒤에 추가로 적은 유의어(~), 반의어(<>), 파생어(>) 처리
    for item in parts[2:]:
        item_str = item.strip()
        if item_str.startswith("~"):
            sims = item_str.replace("~", "", 1).strip()
            for s in sims.split(","):
                if s.strip(): extracted_similar.append(s.strip())
        elif item_str.startswith("<>"):
            ants = item_str.replace("<>", "", 1).strip()
            for a in ants.split(","):
                if a.strip(): extracted_antonyms.append(a.strip())
        elif item_str.startswith(">"):
            ders = item_str.replace(">", "", 1).strip()
            for d in ders.split(","):
                if d.strip(): extracted_derivatives.append(d.strip())

    return {
        "word": word,
        "meaning_pairs": meaning_pairs,
        "similar": ", ".join(dict.fromkeys(extracted_similar)),
        "antonym": ", ".join(dict.fromkeys(extracted_antonyms)),
        "derivative": ", ".join(dict.fromkeys(extracted_derivatives)),
        "grammar": grammar_note
    }

def create_powerpoint():
    vocab_data = load_vocab_file(TXT_PATH)
    print(f"총 {len(vocab_data)}개의 단어 로드 완료. 6단계 PPT 생성 시작!")

    prs = Presentation()
    prs.slide_width = Inches(7.5)
    prs.slide_height = Inches(13.33)
    
    blank_layout = prs.slide_layouts[6]

    BG_COLOR = RGBColor(20, 24, 33)       
    TEXT_WHITE = RGBColor(255, 255, 255)  
    TEXT_GRAY = RGBColor(148, 163, 184)   
    TEXT_THEME = RGBColor(74, 222, 128)   

    for item in vocab_data:
        word = item["word"]
        pairs = item["meaning_pairs"]
        similar = item["similar"]
        derivative = item["derivative"]
        grammar = item["grammar"]

        # -------------------------------------------------------------------------
        # [단계 1] 첫화면: 영단어만 중앙에 크게
        # -------------------------------------------------------------------------
        slide1 = prs.slides.add_slide(blank_layout)
        slide1.background.fill.solid()
        slide1.background.fill.fore_color.rgb = BG_COLOR

        box = slide1.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(6.5), Inches(2.0))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = word
        p.font.size = Pt(72)
        p.font.color.rgb = TEXT_WHITE
        p.font.name = "맑은 고딕"
        p.alignment = PP_ALIGN.CENTER

        # -------------------------------------------------------------------------
        # [단계 2 ~ 3] 뜻 및 뜻=동의어 누적 단계들 생성
        # -------------------------------------------------------------------------
        # 각 뜻(세트)마다 누적 슬라이드 생성
        for idx in range(len(pairs)):
            slide = prs.slides.add_slide(blank_layout)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = BG_COLOR

            # 현재 단계까지 누적된 뜻 문자열 만들기 (좌상단 및 중앙용)
            current_meanings_for_header = []
            current_meanings_for_center = []
            
            for i in range(idx + 1):
                p_item = pairs[i]
                m_str = f"? {p_item['meaning']}"
                if p_item['synonym']:
                    m_str += f" ({p_item['synonym']})"
                current_meanings_for_header.append(m_str)
                current_meanings_for_center.append(m_str)

            # 1) 좌상단 헤더 (단어 + 누적 뜻)
            h_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.5), Inches(3.0))
            tf_h = h_box.text_frame
            tf_h.word_wrap = True
            p_h = tf_h.paragraphs[0]
            p_h.text = f"{word} : " + "\n".join(current_meanings_for_header)
            p_h.font.size = Pt(16)
            p_h.font.color.rgb = TEXT_GRAY
            p_h.font.name = "맑은 고딕"

            # 2) 화면 중앙 뜻 출력
            c_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.0), Inches(6.5), Inches(5.0))
            tf_c = c_box.text_frame
            tf_c.word_wrap = True
            p_c = tf_c.paragraphs[0]
            p_c.text = "\n".join(current_meanings_for_center)
            p_c.font.size = Pt(36)
            p_c.font.color.rgb = TEXT_THEME
            p_c.font.name = "맑은 고딕"
            p_c.alignment = PP_ALIGN.CENTER

        # -------------------------------------------------------------------------
        # [단계 4] 유의어 슬라이드 (유의어가 있는 경우에만 추가)
        # -------------------------------------------------------------------------
        if similar:
            slide_sim = prs.slides.add_slide(blank_layout)
            slide_sim.background.fill.solid()
            slide_sim.background.fill.fore_color.rgb = BG_COLOR

            # 이전 모든 뜻 누적 표시
            all_meanings = [f"? {p['meaning']}" + (f" ({p['synonym']})" if p['synonym'] else "") for p in pairs]
            
            h_box = slide_sim.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.5), Inches(3.0))
            tf_h = h_box.text_frame
            tf_h.word_wrap = True
            p_h = tf_h.paragraphs[0]
            p_h.text = f"{word} : " + "\n".join(all_meanings)
            p_h.font.size = Pt(16)
            p_h.font.color.rgb = TEXT_GRAY
            p_h.font.name = "맑은 고딕"

            c_box = slide_sim.shapes.add_textbox(Inches(0.5), Inches(5.0), Inches(6.5), Inches(5.0))
            tf_c = c_box.text_frame
            tf_c.word_wrap = True
            p_c = tf_c.paragraphs[0]
            p_c.text = f"[유의어]\n{similar}"
            p_c.font.size = Pt(36)
            p_c.font.color.rgb = TEXT_THEME
            p_c.font.name = "맑은 고딕"
            p_c.alignment = PP_ALIGN.CENTER

        # -------------------------------------------------------------------------
        # [단계 5] 파생어 슬라이드 (파생어가 있는 경우에만 추가)
        # -------------------------------------------------------------------------
        if derivative:
            slide_der = prs.slides.add_slide(blank_layout)
            slide_der.background.fill.solid()
            slide_der.background.fill.fore_color.rgb = BG_COLOR

            all_meanings = [f"? {p['meaning']}" + (f" ({p['synonym']})" if p['synonym'] else "") for p in pairs]
            
            h_box = slide_der.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.5), Inches(3.0))
            tf_h = h_box.text_frame
            tf_h.word_wrap = True
            p_h = tf_h.paragraphs[0]
            p_h.text = f"{word} : " + "\n".join(all_meanings)
            p_h.font.size = Pt(16)
            p_h.font.color.rgb = TEXT_GRAY
            p_h.font.name = "맑은 고딕"

            c_box = slide_der.shapes.add_textbox(Inches(0.5), Inches(5.0), Inches(6.5), Inches(5.0))
            tf_c = c_box.text_frame
            tf_c.word_wrap = True
            p_c = tf_c.paragraphs[0]
            p_c.text = f"[파생어]\n{derivative}"
            p_c.font.size = Pt(36)
            p_c.font.color.rgb = TEXT_THEME
            p_c.font.name = "맑은 고딕"
            p_c.alignment = PP_ALIGN.CENTER

        # -------------------------------------------------------------------------
        # [단계 6] 문법 팁 슬라이드 (문법 팁이 있는 경우에만 추가)
        # -------------------------------------------------------------------------
        if grammar:
            slide_gram = prs.slides.add_slide(blank_layout)
            slide_gram.background.fill.solid()
            slide_gram.background.fill.fore_color.rgb = BG_COLOR

            all_meanings = [f"? {p['meaning']}" + (f" ({p['synonym']})" if p['synonym'] else "") for p in pairs]
            
            h_box = slide_gram.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.5), Inches(3.0))
            tf_h = h_box.text_frame
            tf_h.word_wrap = True
            p_h = tf_h.paragraphs[0]
            p_h.text = f"{word} : " + "\n".join(all_meanings)
            p_h.font.size = Pt(16)
            p_h.font.color.rgb = TEXT_GRAY
            p_h.font.name = "맑은 고딕"

            c_box = slide_gram.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(6.5), Inches(6.0))
            tf_c = c_box.text_frame
            tf_c.word_wrap = True
            p_c = tf_c.paragraphs[0]
            p_c.text = f"[문법 팁]\n{grammar}"
            p_c.font.size = Pt(24)
            p_c.font.color.rgb = TEXT_WHITE
            p_c.font.name = "맑은 고딕"
            p_c.alignment = PP_ALIGN.CENTER

    OUTPUT_PPTX_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PPTX_PATH))
    print(f"\n🎉 6단계 순서형 PPT 파일 생성 완료: {OUTPUT_PPTX_PATH.resolve()}")

if __name__ == "__main__":
    create_powerpoint()